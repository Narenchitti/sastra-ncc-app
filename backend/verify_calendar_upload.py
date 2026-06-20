import os
import io
import httpx
import docx
import openpyxl
from app.services import database

async def run_upload_tests():
    print("--- Running Academic Calendar Document Upload & Parse Tests ---")
    
    # 1. Generate temp files
    txt_path = "test_calendar.txt"
    docx_path = "test_calendar.docx"
    xlsx_path = "test_calendar.xlsx"
    
    # Write TXT
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("TXT: Exam Week: July 12 to July 18. Independence Day Holiday: August 15.")
        
    # Write DOCX
    doc = docx.Document()
    doc.add_paragraph("DOCX: Semester Exams scheduled from July 12 to July 18. Working Saturday on August 15.")
    doc.save(docx_path)
    
    # Write XLSX
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NCC Calendar"
    ws.append(["Category", "Holiday details", "Date info"])
    ws.append(["Exams", "Semester Exams", "July 12 to July 18"])
    ws.append(["National Holiday", "Independence Day", "August 15"])
    wb.save(xlsx_path)
    
    # Setup test HTTP client and base URL
    base_url = "http://127.0.0.1:8000/api"
    
    # We will log in as ANO
    # Default ANO credentials from dev environment
    login_payload = {
        "email": "ano@sastra.ncc",
        "password": "12345678"
    }
    
    async with httpx.AsyncClient() as client:
        # Login
        try:
            r = await client.post(f"{base_url}/auth/login", json=login_payload)
            if r.status_code != 200:
                print(f"[FAIL] Login failed: {r.text}")
                return
            res_data = r.json()
            token = res_data.get("accessToken") or res_data.get("access_token")
            if not token:
                print(f"[FAIL] Access token not found in login response: {res_data}")
                return
            headers = {"Authorization": f"Bearer {token}"}
        except Exception as e:
            print(f"[FAIL] Connection failed (make sure backend server is running on port 8000): {e}")
            # Clean up files
            for p in [txt_path, docx_path, xlsx_path]:
                if os.path.exists(p):
                    os.remove(p)
            return

        # Test 1: Upload TXT
        print("\n1. Testing TXT upload...")
        with open(txt_path, "rb") as f:
            files = {"file": (txt_path, f, "text/plain")}
            res = await client.post(f"{base_url}/unit-config/upload-calendar", files=files, headers=headers)
            if res.status_code == 200 and "TXT:" in res.json().get("text", ""):
                print("[PASS] TXT upload and parse passed.")
                print(f"Extracted: {res.json()['text'][:80]}...")
            else:
                print(f"[FAIL] TXT upload failed: {res.status_code} - {res.text}")

        # Test 2: Upload DOCX
        print("\n2. Testing DOCX upload...")
        with open(docx_path, "rb") as f:
            files = {"file": (docx_path, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            res = await client.post(f"{base_url}/unit-config/upload-calendar", files=files, headers=headers)
            if res.status_code == 200 and "DOCX:" in res.json().get("text", ""):
                print("[PASS] DOCX upload and parse passed.")
                print(f"Extracted: {res.json()['text'][:80]}...")
            else:
                print(f"[FAIL] DOCX upload failed: {res.status_code} - {res.text}")

        # Test 3: Upload XLSX
        print("\n3. Testing XLSX upload...")
        with open(xlsx_path, "rb") as f:
            files = {"file": (xlsx_path, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
            res = await client.post(f"{base_url}/unit-config/upload-calendar", files=files, headers=headers)
            if res.status_code == 200 and "Semester Exams" in res.json().get("text", ""):
                print("[PASS] XLSX upload and parse passed.")
                print(f"Extracted: {res.json()['text'][:120]}...")
            else:
                print(f"[FAIL] XLSX upload failed: {res.status_code} - {res.text}")
                
        # Verify db config reflects last save
        print("\n4. Verifying database state...")
        from app.services import sqlite_db
        config = await sqlite_db.get_unit_config()
        if config and "Semester Exams" in (config.get("academic_calendar") or ""):
            print("[PASS] Database unit_config successfully updated with parsed calendar.")
        else:
            print("[FAIL] Database update check failed.")

    # Clean up files
    for p in [txt_path, docx_path, xlsx_path]:
        if os.path.exists(p):
            os.remove(p)
            
if __name__ == "__main__":
    import asyncio
    asyncio.run(run_upload_tests())
